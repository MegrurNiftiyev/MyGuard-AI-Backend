"""
Standalone RETVec+CNN Keras model training & held-out test evaluation script.

Usage:
    python train_model.py
    python -m app.scripts.train_model
"""

import os
import sys
import random
import zipfile
import docx
import pypdf
from pptx import Presentation
import numpy as np

os.environ["TF_USE_LEGACY_KERAS"] = "1"
os.environ["CUDA_VISIBLE_DEVICES"] = "-1"
sys.stdout.reconfigure(encoding='utf-8')

# Ensure project root is in sys.path
BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

SEED = 42
random.seed(SEED)
np.random.seed(SEED)

import tensorflow as tf
tf.random.set_seed(SEED)

from app.ml.cnn.architecture import build_model, LABEL_NAMES
from app.ml.training.dataset import encode_labels
from app.ml.training.train import get_class_weights
from app.ml.preprocessing.chunking import chunk_text

HELDOUT_TEST_FILES = {
    "benign": [
        "09_resmi_mektub_temiz.docx",
        "10_iclas_protokolu_temiz.docx",
        "Monthly Financial Expense Report.pdf",
        "11_ezamiyye_emri_temiz.docx",
        "19_sifaris_senedi_temiz.docx"
    ],
    "injection": [
        "01_Aylıq_Fəaliyyət_Hesabatı.docx",
        "16_ezamiyye_xercleri_injection_gizli.docx",
        "19_sifaris_senedi_problem.docx",
        "23_bank_zemanet_mektubu_injection_context_hijack.docx",
        "24_qebul_tehvil_akti_injection.docx"
    ]
}


def extract_pptx(file_path: str) -> str:
    """Extract slide paragraph text and notes text from PPTX files using python-pptx."""
    try:
        prs = Presentation(file_path)
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in para.runs)
                        if line.strip():
                            parts.append(line.strip())
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                note = slide.notes_slide.notes_text_frame.text
                if note.strip():
                    parts.append(note.strip())
        return "\n".join(parts)
    except Exception as e:
        print(f"Warning reading PPTX {file_path}: {e}")
        return ""


def extract_text(file_path: str) -> str:
    """Extract raw text from supported document formats (.docx, .pptx, .pdf, .zip, .txt)."""
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    try:
        if ext == ".docx":
            doc = docx.Document(file_path)
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            parts.append(cell.text.strip())
            text = "\n".join(parts)
        elif ext == ".pptx":
            text = extract_pptx(file_path)
        elif ext == ".pdf":
            reader = pypdf.PdfReader(file_path)
            parts = []
            for i, page in enumerate(reader.pages):
                if i >= 20:
                    break
                try:
                    t = page.extract_text()
                    if t:
                        parts.append(t.strip())
                except Exception:
                    continue
            text = "\n".join(parts)
        elif ext == ".zip":
            parts = []
            with zipfile.ZipFile(file_path, 'r') as z:
                for name in z.namelist():
                    if name.endswith('.docx'):
                        tmp_path = os.path.join(os.path.dirname(file_path), "_tmp_extracted.docx")
                        with open(tmp_path, "wb") as f_out:
                            f_out.write(z.read(name))
                        sub_text = extract_text(tmp_path)
                        if os.path.exists(tmp_path):
                            os.remove(tmp_path)
                        parts.append(sub_text)
                    elif name.endswith('.pptx'):
                        tmp_path = os.path.join(os.path.dirname(file_path), "_tmp_extracted.pptx")
                        with open(tmp_path, "wb") as f_out:
                            f_out.write(z.read(name))
                        sub_text = extract_text(tmp_path)
                        if os.path.exists(tmp_path):
                            os.remove(tmp_path)
                        parts.append(sub_text)
                    elif name.endswith('.txt'):
                        parts.append(z.read(name).decode('utf-8', errors='ignore'))
            text = "\n".join(parts)
        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        else:
            print(f"Skipping unsupported file extension {ext} for {file_path}")
            return ""
    except Exception as e:
        print(f"Warning reading {file_path}: {e}")
    return text.strip()


def split_documents(doc_ids: list[str], val_ratio: float = 0.15, seed: int = 42) -> tuple[set[str], set[str]]:
    """Perform a document-level split of source document IDs into train and validation sets."""
    rng = random.Random(seed)
    unique_ids = list(dict.fromkeys(doc_ids))
    rng.shuffle(unique_ids)
    n_val = max(1, int(len(unique_ids) * val_ratio))
    val_ids = set(unique_ids[:n_val])
    train_ids = set(unique_ids[n_val:])
    return train_ids, val_ids


def load_real_dataset(raw_dir: str):
    all_chunks = []  # [(doc_id, text_chunk, label)]
    all_doc_ids = []
    test_docs = []

    # Define folder mapping: (folder_path, default_category)
    folders_to_scan = [
        (os.path.join(raw_dir, "benign"), "benign"),
        (os.path.join(raw_dir, "injection"), "injection"),
    ]

    downloaded_dir = os.path.join(raw_dir, "downloaded")
    if os.path.exists(downloaded_dir):
        for root, dirs, files in os.walk(downloaded_dir):
            if files:
                folders_to_scan.append((root, "benign"))

    scanned_file_counts = {}

    for cat_dir, category in folders_to_scan:
        if not os.path.exists(cat_dir):
            continue

        heldout_list = HELDOUT_TEST_FILES.get(category, [])
        label_str = "safe" if category == "benign" else "injection"
        dir_key = os.path.relpath(cat_dir, raw_dir)
        scanned_file_counts[dir_key] = scanned_file_counts.get(dir_key, 0)

        for fname in os.listdir(cat_dir):
            fpath = os.path.join(cat_dir, fname)
            if not os.path.isfile(fpath):
                continue

            extracted = extract_text(fpath)
            if not extracted:
                continue

            scanned_file_counts[dir_key] += 1
            doc_id = os.path.relpath(fpath, raw_dir)

            if fname in heldout_list:
                test_docs.append({
                    "filename": fname,
                    "category": category,
                    "expected_label": label_str,
                    "text": extracted
                })
            else:
                ext = os.path.splitext(fname)[1].lower()
                file_chunks = []
                # Check for docx paragraph-level white font / hidden text
                docx_inj_lines = set()
                if ext == ".docx":
                    try:
                        doc = docx.Document(fpath)
                        for p in doc.paragraphs:
                            ptxt = p.text.strip()
                            if not ptxt:
                                continue
                            is_p_white = False
                            for r in p.runs:
                                if r.font.color and r.font.color.rgb and str(r.font.color.rgb).upper() in ("FFFFFF", "FFF"):
                                    is_p_white = True
                                    break
                                if r._r.rPr is not None and r._r.rPr.find(docx.oxml.ns.qn("w:vanish")) is not None:
                                    is_p_white = True
                                    break
                            if is_p_white:
                                docx_inj_lines.add(ptxt)
                    except Exception:
                        pass

                MANUAL_STEALTHY_LABELS = {
                    "23_bank_zemanet_mektubu_injection_context_hijack.docx": "450,000 azn",
                    "24_qebul_tehvil_akti_injection.docx": "administrative correction",
                    "25_sigorta_polisi_injection.docx": "franchise override protocol",
                    "26_emek_muqavilesi_inject.docx": "critical parser override",
                    "27_vekaletname_inject.docx": "müvəqqəti hüquqi düzəliş",
                    "29_bank_rekvizit_deyisikligi_injection.docx": "sistem təlimatı",
                    "31_tecili_odenis_memo_inject.docx": "suni intelleke",
                    "32_hosting_deyisiklik_bildirisi_inject.docx": "audit id-7741",
                    "33_elave_is_sifarisi_inject.docx": "audit id-112",
                    "34_distributor_endirim_teklifi_inject.docx": "audit id-267",
                    "38_icare_muqavilesi_injection.docx": "tərəfdaşlıq ianəsi",
                    "39_dasima_xidmeti_muqavilesi_inject.docx": "<last update>"
                }

                lines = [l.strip() for l in extracted.split("\n") if l.strip()]
                for line in lines:
                    is_inj_line = False
                    if category == "injection":
                        low = line.lower()
                        if fname in MANUAL_STEALTHY_LABELS:
                            if MANUAL_STEALTHY_LABELS[fname] in low:
                                is_inj_line = True
                        else:
                            low = line.lower()
                            if line in docx_inj_lines or any(kw in low for kw in [
                                "prompt", "system", "yuxarida", "mene", "ignore", "override",
                                "@", "//", "#", "||", "^^", "***", "&&", "<system", "[system",
                                "internal system update", "forget", "unrestricted"
                            ]):
                                is_inj_line = True

                    lbl = "injection" if (category == "injection" and is_inj_line) else "safe"

                    words = line.split()
                    if len(words) <= 60:
                        file_chunks.append((line, lbl))
                    else:
                        for c in chunk_text(line):
                            file_chunks.append((c, lbl))

                # Cap per-document safe chunks so long PDFs don't dominate dataset (max 15 safe chunks per doc)
                if len(file_chunks) > 15:
                    inj_chunks = [c for c in file_chunks if c[1] == "injection"]
                    safe_chunks = [c for c in file_chunks if c[1] == "safe"]
                    needed_safe = max(5, 15 - len(inj_chunks))
                    step = max(1, len(safe_chunks) // needed_safe) if safe_chunks else 1
                    file_chunks = inj_chunks + (safe_chunks[::step][:needed_safe] if safe_chunks else [])

                for text_chunk, lbl in file_chunks:
                    all_chunks.append((doc_id, text_chunk, lbl))
                    all_doc_ids.append(doc_id)

    # Document-level split
    train_doc_ids, val_doc_ids = split_documents(all_doc_ids, val_ratio=0.15, seed=SEED)

    train_tuples = [c for c in all_chunks if c[0] in train_doc_ids]
    val_tuples = [c for c in all_chunks if c[0] in val_doc_ids]

    # Oversample injection training tuples so model learns injection patterns properly
    train_inj_tuples = [t for t in train_tuples if t[2] == "injection"]
    train_safe_tuples = [t for t in train_tuples if t[2] == "safe"]

    if train_inj_tuples and len(train_safe_tuples) > 0:
        multiplier = max(1, (len(train_safe_tuples) // 3) // len(train_inj_tuples))
        train_inj_oversampled = train_inj_tuples * multiplier
        train_tuples = train_safe_tuples + train_inj_oversampled

    # Thorough random shuffling across all sources, classes, and languages
    rng = random.Random(SEED)
    rng.shuffle(train_tuples)
    rng.shuffle(val_tuples)

    train_texts = [t[1] for t in train_tuples]
    train_labels = [t[2] for t in train_tuples]

    val_texts = [t[1] for t in val_tuples]
    val_labels = [t[2] for t in val_tuples]

    print("Scanned files count per folder:")
    for folder_rel, count in scanned_file_counts.items():
        print(f"  - {folder_rel}: {count} valid documents")

    print(f"Document-level split: {len(train_doc_ids)} train docs ({len(train_texts)} chunks), {len(val_doc_ids)} val docs ({len(val_texts)} chunks)")

    return (train_texts, train_labels), (val_texts, val_labels), test_docs


def main():
    raw_dir = os.path.join(BASE_DIR, "data", "raw")
    print("Reading document dataset from data/raw...")

    (train_texts, train_labels), (val_texts, val_labels), test_docs = load_real_dataset(raw_dir)

    print(f"\n--- Dataset Loading Summary ---")
    print(f"Training text chunks extracted: {len(train_texts)}")
    print(f"  - Safe (Benign) train chunks: {train_labels.count('safe')}")
    print(f"  - Injection train chunks: {train_labels.count('injection')}")
    print(f"Validation text chunks extracted: {len(val_texts)}")
    print(f"  - Safe (Benign) val chunks: {val_labels.count('safe')}")
    print(f"  - Injection val chunks: {val_labels.count('injection')}")
    print(f"Held-out Test Files reserved: {len(test_docs)}")
    for td in test_docs:
        print(f"  * [{td['category'].upper()}] {td['filename']} ({len(td['text'])} chars)")

    X_train = np.array([[t] for t in train_texts])
    Y_train_label = encode_labels(train_labels)

    X_val = np.array([[t] for t in val_texts])
    Y_val_label = encode_labels(val_labels)

    class_weights_dict = get_class_weights(Y_train_label)
    sample_weights_label = np.array([class_weights_dict[int(np.argmax(y))] for y in Y_train_label], dtype=np.float32)

    print("\nBuilding RETVec + CNN Keras Classification Model...")
    model = build_model(sequence_length=128)
    model.summary()

    print("\nStarting Keras Model Training (5 Epochs, batch_size=128, document-level validation)...", flush=True)
    history = model.fit(
        X_train,
        Y_train_label,
        epochs=5,
        batch_size=128,
        validation_data=(X_val, Y_val_label),
        sample_weight=sample_weights_label,
        verbose=1
    )

    models_dir = os.path.join(BASE_DIR, "data", "models")
    os.makedirs(models_dir, exist_ok=True)
    keras_model_path = os.path.join(models_dir, "retvec_cnn_model.keras")
    
    print(f"\nSaving trained model to .keras file at:\n  {keras_model_path}")
    model.save(keras_model_path)

    cache_dir = os.path.join(BASE_DIR, "data", "cache")
    os.makedirs(cache_dir, exist_ok=True)
    model.save(os.path.join(cache_dir, "active_model.keras"))

    print("\n==========================================")
    print("HELD-OUT TEST FILES INFERENCE & EVALUATION")
    print("==========================================")

    correct_predictions = 0
    test_results = []

    for td in test_docs:
        raw_text = td["text"]
        lines = [l.strip() for l in raw_text.split("\n") if l.strip()]
        chunks = []
        for line in lines:
            words = line.split()
            if len(words) <= 60:
                chunks.append(line)
            else:
                chunks.extend(chunk_text(line))

        chunk_inputs = np.array([[c] for c in chunks])
        
        preds = model.predict(chunk_inputs, verbose=0)
        label_preds = preds if isinstance(preds, np.ndarray) and preds.ndim == 2 else preds[0]

        worst_chunk_idx = label_preds[:, 2].argmax()
        max_injection_prob = float(label_preds[worst_chunk_idx, 2])
        max_inj_line = chunks[worst_chunk_idx] if chunks else ""

        avg_probs = np.mean(label_preds, axis=0)

        HIGH_CONF_THRESHOLD = 0.85
        CORROBORATION_THRESHOLD = 0.60
        MIN_CORROBORATING_CHUNKS = 2

        injection_probs = [float(p) for p in label_preds[:, 2]]
        
        predicted_label = "safe"
        high_conf = [p for p in injection_probs if p >= HIGH_CONF_THRESHOLD]
        if high_conf:
            predicted_label = "injection"
        else:
            corroborating = [p for p in injection_probs if p >= CORROBORATION_THRESHOLD]
            if len(corroborating) >= MIN_CORROBORATING_CHUNKS:
                predicted_label = "injection"

        is_correct = (predicted_label == td["expected_label"])
        if is_correct:
            correct_predictions += 1

        test_results.append({
            "filename": td["filename"],
            "expected": td["expected_label"],
            "predicted": predicted_label,
            "is_correct": is_correct,
            "prob_safe": float(avg_probs[0]),
            "prob_suspicious": float(avg_probs[1]),
            "prob_injection": float(avg_probs[2]),
            "max_chunk_injection": float(max_injection_prob),
            "max_inj_snippet": max_inj_line[:60]
        })

        status = "PASSED ✓" if is_correct else "FAILED ✗"
        print(f"File: {td['filename']}")
        print(f"  Expected: {td['expected_label']} | Predicted: {predicted_label} [{status}]")
        print(f"  Max Injection Prob: {max_injection_prob:.2%} | Snippet: {max_inj_line[:70]!r}\n")

    accuracy = (correct_predictions / len(test_docs)) * 100 if test_docs else 0.0
    print(f"Final Held-Out Test Accuracy: {accuracy:.2f}% ({correct_predictions}/{len(test_docs)})")


if __name__ == "__main__":
    main()
