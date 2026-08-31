import os
import sys
import zipfile
import docx
import pypdf

sys.stdout.reconfigure(encoding='utf-8')

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    try:
        if ext == ".docx":
            doc = docx.Document(file_path)
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            parts.append(cell.text.strip())
            text = "\n".join(parts)
        elif ext == ".pdf":
            reader = pypdf.PdfReader(file_path)
            parts = []
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    parts.append(t.strip())
            text = "\n".join(parts)
        elif ext == ".zip":
            parts = []
            with zipfile.ZipFile(file_path, 'r') as z:
                for name in z.namelist():
                    if name.endswith('.txt'):
                        parts.append(z.read(name).decode('utf-8', errors='ignore'))
            text = "\n".join(parts)
        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
    except Exception as e:
        print(f"Error reading {file_path}: {e}")
    return text.strip()

raw_dir = r"c:\Users\megru\Desktop\Programlar\Github\MyGurad-IDDA-Final_project\Ai-Models\data\raw"
for category in ["benign", "injection"]:
    cat_dir = os.path.join(raw_dir, category)
    print(f"=== {category.upper()} FILES ===")
    files = os.listdir(cat_dir)
    print(f"Total files: {len(files)}")
    valid_count = 0
    for f in files:
        fp = os.path.join(cat_dir, f)
        txt = extract_text(fp)
        if txt:
            valid_count += 1
            print(f"  [OK] {f} -> {len(txt)} chars | Snippet: {txt[:80]!r}")
        else:
            print(f"  [EMPTY/FAIL] {f}")
    print(f"Valid extracted: {valid_count}/{len(files)}\n")
